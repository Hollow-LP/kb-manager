"""生成 seed_docs/ 下的示例知识文档（PDF / DOCX / MD / TXT）。

内容为虚构公司的制度文档，用于演示导入解析、切片与四维数据权限。
    uv run python scripts/gen_seed_docs.py
"""

from pathlib import Path

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parents[1] / "seed_docs"

DOCS: dict[str, dict] = {
    "员工报销制度.docx": {
        "type": "docx",
        "sections": [
            ("一、适用范围", [
                "本制度适用于公司全体在职员工因公出差、业务招待、办公采购等产生的费用报销。",
                "实习生的差旅费用按正式员工标准的百分之八十执行，须由导师代为提交报销单。",
            ]),
            ("二、报销流程", [
                "1. 员工应在费用发生后十五个工作日内提交报销申请，先由直属上级审批，再由财务中心复核。",
                "2. 单笔金额超过五千元需部门负责人与财务负责人双签。",
                "3. 财务复核通过后，报销款在每个工作日的下午统一打款至员工工资卡。",
            ]),
            ("三、票据要求", [
                "报销须提供合规发票原件或电子发票 PDF，抬头必须为公司全称。",
                "餐费、交通费须附消费清单；超过三百元的打车费须说明事由。",
                "发票遗失的，经办人须填写《发票遗失说明》并由部门负责人签字。",
            ]),
            ("四、常见问题", [
                "问：报销单被财务退回怎么办？答：查看退回原因，补正后在原单上重新提交即可，无需重新填写。",
                "问：跨年的发票能否报销？答：上一年度十二月及以后的发票最迟在次年一月十五日前提交。",
            ]),
        ],
    },
    "IT安全规范.pdf": {
        "type": "pdf",
        "sections": [
            ("第一章 总则", [
                "为保障公司信息系统与数据安全，所有使用公司网络的员工须遵守本规范。",
                "各部门负责人为部门信息安全第一责任人。",
            ]),
            ("第二章 账号与密码", [
                "员工账号须使用长度不少于十二位、包含大小写字母与数字的强密码，每九十天更换一次。",
                "禁止多人共用同一账号，员工离职时须在当日内回收全部系统权限。",
                "凡涉及生产环境的账号须开启双因素认证。",
            ]),
            ("第三章 终端安全", [
                "办公电脑须开启全盘加密并安装公司统一的安全客户端。",
                "禁止在办公电脑上安装未经审批的远程控制软件。",
                "U 盘等移动存储设备须先在隔离区查杀病毒后方可使用。",
            ]),
            ("第四章 数据外发", [
                "包含客户信息的文件外发须经信息安全组审批，并通过公司指定的加密通道传输。",
                "严禁通过个人网盘、邮箱转发公司源代码与客户数据库导出文件。",
            ]),
        ],
    },
    "财务审批权限手册.docx": {
        "type": "docx",
        "sections": [
            ("一、审批层级", [
                "部门负责人可审批单笔一万元以内的费用；财务负责人可审批五万元以内；超过五万元须由总经理审批。",
                "对外付款须同时具备合同编号与验收单，否则财务有权拒绝受理。",
            ]),
            ("二、预算管理", [
                "各部门每月费用不得超过预算额度，超预算部分须提前提交预算调整申请。",
                "预算调整超过原有预算百分之二十的，须经财务委员会审议。",
            ]),
            ("三、资金安全", [
                "银行 U 盾由财务中心两人分开保管，付款操作须双人复核。",
                "大额付款前须通过电话回拨核实收款账户。",
            ]),
        ],
    },
    "新员工入职指南.md": {
        "type": "md",
        "sections": [
            ("入职第一周", [
                "1. 报到当天到行政中心领取工牌、电脑与门禁卡。",
                "2. 登录 OKR 系统完成个人信息填写，并加入部门企业微信群。",
                "3. 参加新人培训，内容包括公司制度、信息安全与产品概览。",
            ]),
            ("办公环境", [
                "工位可提前在企业微信的行政服务号中预约。会议室通过会议室系统预订，使用后请按时释放。",
                "公司提供午餐补贴，每月随工资发放。",
            ]),
            ("试用期须知", [
                "试用期为三个月，转正需提交转正申请并由直属上级与 HR 面谈。",
                "试用期内请假仍需按考勤制度提交申请。",
            ]),
        ],
    },
    "数据合规红线.md": {
        "type": "md",
        "sections": [
            ("红线清单", [
                "- 禁止将真实用户数据导入非公司授权的大模型服务。",
                "- 禁止在公开代码仓库中提交任何密钥、证书或连接串。",
                "- 禁止擅自将生产数据库副本复制到个人设备。",
            ]),
            ("违规处理", [
                "违反数据合规红线的人员，视情节给予书面警告、绩效降级直至解除劳动合同，并保留追究法律责任的权利。",
                "发现违规线索可向合规邮箱 compliance@example.com 匿名举报。",
            ]),
        ],
    },
    "差旅标准说明.txt": {
        "type": "txt",
        "sections": [
            ("一、交通标准", [
                "普通员工出差可乘坐高铁二等座或经济舱；部门负责人可乘坐高铁一等座；总监及以上可乘坐公务舱。",
                "城市内交通优先使用公共交通，特殊情况可乘坐网约车，单次不得超过一百元。",
            ]),
            ("二、住宿标准", [
                "一线城市普通员工每晚不超过六百元，其他城市不超过四百元；部门负责人一线城市不超过九百元。",
                "住宿须提供酒店发票，超过标准的差额由个人承担。",
            ]),
            ("三、出差补贴", [
                "出差期间按每人每天一百五十元发放餐补，出差结束后随报销单一并提交。",
                "周末出差可申请调休，调休须在返回后一个月内使用。",
            ]),
        ],
    },
    "无权限配置的测试文档.txt": {
        "type": "txt",
        "sections": [
            ("测试说明", [
                "本文档用于验证知识单元的默认无权限状态：未配置任何数据权限实体时，所有用户（包括知识管理员）在 AI 问答中都无法访问其内容。",
                "文件包含一段虚构的年度经营数据：本年度营业收入三万万元，同比增长百分之十二。",
            ]),
        ],
    },
    "年假计算规则.docx": {
        "type": "docx",
        "sections": [
            ("一、年假天数", [
                "员工累计工作满一年不满十年的，年假五天；满十年不满二十年的，年假十天；满二十年的，年假十五天。",
                "新员工入职当年按剩余日历天数折算年假，折算后不足一天的按一天计。",
            ]),
            ("二、年假申请", [
                "年假须提前三个工作日提交，连续请假超过五天的须经部门负责人审批。",
                "当年未休完的年假最多结转至次年三月三十一日。",
            ]),
            ("三、未休年假补偿", [
                "因工作原因未能安排休年假的，经 HR 确认后按日工资的百分之二百支付补偿。",
            ]),
        ],
    },
}


def build_docx(path: Path, sections: list[tuple[str, list[str]]]) -> None:
    doc = Document()
    doc.add_heading(path.stem, level=0)
    for heading, paras in sections:
        doc.add_heading(heading, level=1)
        for p in paras:
            doc.add_paragraph(p)
    doc.save(path)


def build_pdf(path: Path, sections: list[tuple[str, list[str]]]) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    title = ParagraphStyle("title", fontName="STSong-Light", fontSize=18, spaceAfter=12)
    h = ParagraphStyle("h", fontName="STSong-Light", fontSize=14, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("body", fontName="STSong-Light", fontSize=11, leading=18)
    story: list = [Paragraph(path.stem, title)]
    for heading, paras in sections:
        story.append(Paragraph(heading, h))
        for p in paras:
            story.append(Paragraph(p, body))
            story.append(Spacer(1, 4))
    raw = path.with_suffix(".raw.pdf")
    SimpleDocTemplate(str(raw), pagesize=A4, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54).build(story)
    # 设置文档标题元数据
    reader = PdfReader(str(raw))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.add_metadata({"/Title": path.stem, "/Author": "行政中心"})
    with open(path, "wb") as f:
        writer.write(f)
    raw.unlink()


def build_text(path: Path, sections: list[tuple[str, list[str]]]) -> None:
    lines: list[str] = [path.stem, ""]
    for heading, paras in sections:
        lines.append(heading)
        lines.extend(paras)
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, spec in DOCS.items():
        path = OUT / name
        kind = spec["type"]
        if kind == "docx":
            build_docx(path, spec["sections"])
        elif kind == "pdf":
            build_pdf(path, spec["sections"])
        else:
            build_text(path, spec["sections"])
        print(f"{kind:5s} {name:28s} {path.stat().st_size:>7d} bytes")


if __name__ == "__main__":
    main()
